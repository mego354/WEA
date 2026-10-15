from decimal import Decimal
import io
import openpyxl
from django.test import TestCase, Client
from django.urls import reverse
from django.contrib.auth import get_user_model
from contacts.models import Contact
from sales.models import SalesOrder, Collection, CustomerAccount

User = get_user_model()


class SalesExcelExportTestCase(TestCase):
    def setUp(self):
        self.user = User.objects.create_superuser(
            email='salesadmin@weapharma.com',
            username='salesadmin@weapharma.com',
            password='Password123!',
            first_name='Sales',
            last_name='Admin'
        )
        self.client = Client()
        self.client.login(username='salesadmin@weapharma.com', password='Password123!')

        self.contact = Contact.objects.create(
            serial_num=8801,
            name='صيدلية الشفاء التخصصية',
            is_customer=True,
            phone='01234567890'
        )
        self.customer_account = CustomerAccount.objects.create(
            customer=self.contact,
            credit_limit=Decimal('100000.00'),
            current_balance=Decimal('25000.00')
        )
        from inventory.models import Warehouse
        from treasury.models import Vault

        self.warehouse = Warehouse.objects.create(name='مخزن مبيعات رئيسي', code='WH-SALES')
        self.vault = Vault.objects.create(name='خزينة مبيعات', vault_type='cash')

        self.order = SalesOrder.objects.create(
            customer=self.contact,
            source_warehouse=self.warehouse,
            status='pending',
            grand_total=Decimal('4800.00'),
            created_by=self.user
        )
        self.collection = Collection.objects.create(
            customer=self.contact,
            vault=self.vault,
            amount=Decimal('4800.00'),
            payment_method='cash',
            created_by=self.user
        )

    def test_sales_orders_export_excel(self):
        response = self.client.get(reverse('sales_orders_export'))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Type'], 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        self.assertIn('.xlsx', response['Content-Disposition'])

        wb = openpyxl.load_workbook(io.BytesIO(response.content))
        ws = wb.active
        self.assertTrue(ws.sheet_view.rightToLeft)
        found = any(cell.value == self.order.order_number for row in ws.iter_rows() for cell in row)
        self.assertTrue(found)

    def test_collections_export_excel(self):
        response = self.client.get(reverse('collections_export'))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Type'], 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        self.assertIn('.xlsx', response['Content-Disposition'])

        wb = openpyxl.load_workbook(io.BytesIO(response.content))
        ws = wb.active
        self.assertTrue(ws.sheet_view.rightToLeft)
        found = any('صيدلية الشفاء' in str(cell.value) for row in ws.iter_rows() for cell in row)
        self.assertTrue(found)
