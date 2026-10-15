from decimal import Decimal
from django.test import TestCase, Client
from django.urls import reverse
from django.contrib.auth import get_user_model
from catalog.models import Product, ProductCategory

User = get_user_model()


class CatalogModelsAndViewsTestCase(TestCase):
    def setUp(self):
        self.user = User.objects.create_superuser(
            email='testuser@weapharma.com',
            username='testuser@weapharma.com',
            password='Password123!',
            first_name='Admin',
            last_name='User'
        )
        self.client = Client()
        self.client.login(username='testuser@weapharma.com', password='Password123!')

        self.category = ProductCategory.objects.create(
            name='مكملات غذائية',
            code='SUPP',
            description='قسم المكملات الغذائية والفيتامينات'
        )

        self.product = Product.objects.create(
            name='Geron',
            sku_code='GER-001',
            barcode='6223001234567',
            category=self.category,
            retail_price=Decimal('18.80'),
            tax_percentage=Decimal('0.00'),
            default_purchase_cost=Decimal('15.00'),
            packing_boxes='10',
            packing_cartons='50',
            has_volume_discounts=True,
            sellable=True,
            unit='علبة',
            description='كبسولات جيرون مكمل غذائي'
        )

    def test_product_model_fields_and_methods(self):
        self.assertEqual(str(self.product), 'Geron (GER-001)')
        self.assertEqual(self.product.barcode, '6223001234567')
        self.assertEqual(self.product.packing_boxes, '10')
        self.assertEqual(self.product.packing_cartons, '50')
        self.assertTrue(self.product.has_volume_discounts)
        self.assertTrue(self.product.sellable)
        self.assertIn('18.80', self.product.price_display)

    def test_category_model_str(self):
        self.assertEqual(str(self.category), 'مكملات غذائية')

    def test_products_list_view(self):
        response = self.client.get(reverse('products_list'))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Geron')
        self.assertContains(response, 'GER-001')

    def test_products_list_filtering(self):
        # Search filter
        resp = self.client.get(reverse('products_list'), {'search': 'Geron'})
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, 'GER-001')

        resp_none = self.client.get(reverse('products_list'), {'search': 'NonExistentXYZ'})
        self.assertEqual(resp_none.status_code, 200)
        self.assertNotContains(resp_none, 'GER-001')

        # Category filter
        resp_cat = self.client.get(reverse('products_list'), {'category': self.category.id})
        self.assertEqual(resp_cat.status_code, 200)
        self.assertContains(resp_cat, 'Geron')

        # Sellable filter
        resp_sellable = self.client.get(reverse('products_list'), {'sellable': 'true'})
        self.assertEqual(resp_sellable.status_code, 200)
        self.assertContains(resp_sellable, 'Geron')

        # Price range filter
        resp_price = self.client.get(reverse('products_list'), {'min_price': '10', 'max_price': '30'})
        self.assertEqual(resp_price.status_code, 200)
        self.assertContains(resp_price, 'Geron')

    def test_product_create_view_get(self):
        response = self.client.get(reverse('product_create'))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'انشئ منتج')

    def test_product_create_view_post(self):
        data = {
            'name': 'Legatonexx Spray',
            'sku_code': 'LEG-002',
            'barcode': '6224009876543',
            'category': self.category.id,
            'retail_price': '45.00',
            'tax_percentage': '14.00',
            'default_purchase_cost': '30.00',
            'packing_boxes': '1',
            'packing_cartons': '24',
            'has_volume_discounts': 'on',
            'sellable': 'on',
            'unit': 'زجاجة',
            'description': 'بخاخ موضعي',
        }
        response = self.client.post(reverse('product_create'), data, follow=True)
        self.assertEqual(response.status_code, 200)
        self.assertTrue(Product.objects.filter(sku_code='LEG-002').exists())
        created_p = Product.objects.get(sku_code='LEG-002')
        self.assertEqual(created_p.name, 'Legatonexx Spray')
        self.assertEqual(created_p.packing_cartons, '24')

    def test_product_edit_view(self):
        edit_url = reverse('product_edit', args=[self.product.id])
        resp = self.client.get(edit_url)
        self.assertEqual(resp.status_code, 200)

        data = {
            'name': 'Geron Updated',
            'sku_code': 'GER-001',
            'barcode': '6223001234567',
            'category': self.category.id,
            'retail_price': '22.50',
            'tax_percentage': '0.00',
            'default_purchase_cost': '18.00',
            'packing_boxes': '12',
            'packing_cartons': '60',
            'unit': 'علبة',
            'description': 'محدث',
        }
        resp_post = self.client.post(edit_url, data, follow=True)
        self.assertEqual(resp_post.status_code, 200)
        self.product.refresh_from_db()
        self.assertEqual(self.product.name, 'Geron Updated')
        self.assertEqual(self.product.retail_price, Decimal('22.50'))

    def test_product_detail_view(self):
        detail_url = reverse('product_detail', args=[self.product.id])
        response = self.client.get(detail_url)
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Geron')
        self.assertContains(response, 'GER-001')

    def test_product_delete_view(self):
        delete_url = reverse('product_delete', args=[self.product.id])
        response = self.client.post(delete_url, follow=True)
        self.assertEqual(response.status_code, 200)
        self.assertFalse(Product.objects.filter(id=self.product.id).exists())

    def test_products_export_excel(self):
        import io
        import openpyxl
        response = self.client.get(reverse('products_export'))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Type'], 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        self.assertIn('.xlsx', response['Content-Disposition'])
        
        # Verify it is a valid openpyxl workbook
        wb = openpyxl.load_workbook(io.BytesIO(response.content))
        ws = wb.active
        self.assertTrue(ws.sheet_view.rightToLeft)
        # Verify Geron is in the sheet
        found = any(cell.value == 'Geron' for row in ws.iter_rows() for cell in row)
        self.assertTrue(found)

    def test_product_import_template(self):
        import io
        import openpyxl
        response = self.client.get(reverse('product_import_template'))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Type'], 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        wb = openpyxl.load_workbook(io.BytesIO(response.content))
        ws = wb.active
        self.assertEqual(ws.cell(row=1, column=1).value, 'كود المنتج (SKU)')

    def test_product_import_workflow(self):
        import io
        import openpyxl
        from django.core.files.uploadedfile import SimpleUploadedFile
        # 1. GET page
        get_resp = self.client.get(reverse('product_import'))
        self.assertEqual(get_resp.status_code, 200)
        self.assertContains(get_resp, 'استيراد المنتجات')

        # 2. Create workbook in memory matching product columns
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(['كود المنتج (SKU)', 'اسم المنتج', 'التصنيف', 'سعر البيع للجمهور', 'سعر تكلفة الشراء', 'نسبة الضريبة', 'متاح للبيع', 'الباركود', 'الوحدة', 'الوصف'])
        ws.append(['PAN-COLD-01', 'بانادول كولد', 'مكملات غذائية', 45.0, 32.5, 0.0, 'نعم', '6221234567890', 'علبة', 'مضاد نزلات البرد'])
        out = io.BytesIO()
        wb.save(out)
        out.seek(0)

        excel_file = SimpleUploadedFile(
            'products_test.xlsx',
            out.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )

        # 3. POST preview
        preview_resp = self.client.post(reverse('product_import'), {
            'action': 'preview',
            'excel_file': excel_file,
        })
        self.assertEqual(preview_resp.status_code, 200)
        self.assertContains(preview_resp, 'PAN-COLD-01')
        self.assertContains(preview_resp, 'اختر طريقة تنفيذ الاستيراد')

        # 4. POST execute import
        execute_resp = self.client.post(reverse('product_import'), {
            'action': 'execute',
            'import_mode': 'create_and_update',
        }, follow=True)
        self.assertEqual(execute_resp.status_code, 200)
        self.assertTrue(Product.objects.filter(sku_code='PAN-COLD-01').exists())
        imported_p = Product.objects.get(sku_code='PAN-COLD-01')
        self.assertEqual(imported_p.name, 'بانادول كولد')
        self.assertEqual(imported_p.retail_price, Decimal('45.00'))

    def test_categories_list_and_create(self):
        cat_url = reverse('categories_list')
        response = self.client.get(cat_url)
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'مكملات غذائية')

        # POST new category
        data = {
            'name': 'مستحضرات تجميل',
            'code': 'COSM',
            'description': 'عناية بالبشرة'
        }
        post_resp = self.client.post(cat_url, data, follow=True)
        self.assertEqual(post_resp.status_code, 200)
        self.assertTrue(ProductCategory.objects.filter(code='COSM').exists())

    def test_reference_url_aliases(self):
        # /products/
        resp = self.client.get('/products/')
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, 'Geron')

        # /product-categories/
        resp_cats = self.client.get('/product-categories/')
        self.assertEqual(resp_cats.status_code, 200)
        self.assertContains(resp_cats, 'مكملات غذائية')

    def test_category_edit_and_delete(self):
        # Category edit
        edit_url = reverse('category_edit', args=[self.category.id])
        resp = self.client.post(edit_url, {
            'name': 'مكملات وفيتامينات محدثة',
            'code': 'SUPP-NEW',
            'description': 'وصف محدث'
        }, follow=True)
        self.assertEqual(resp.status_code, 200)
        self.category.refresh_from_db()
        self.assertEqual(self.category.name, 'مكملات وفيتامينات محدثة')
        self.assertEqual(self.category.code, 'SUPP-NEW')

        # Category delete protected when products exist
        del_url = reverse('category_delete', args=[self.category.id])
        del_resp = self.client.post(del_url, follow=True)
        self.assertEqual(del_resp.status_code, 200)
        self.assertTrue(ProductCategory.objects.filter(id=self.category.id).exists())

        # Create empty category and delete
        empty_cat = ProductCategory.objects.create(name='تصنيف مؤقت', code='TEMP')
        del_empty_resp = self.client.post(reverse('category_delete', args=[empty_cat.id]), follow=True)
        self.assertEqual(del_empty_resp.status_code, 200)
        self.assertFalse(ProductCategory.objects.filter(id=empty_cat.id).exists())

    def test_volume_discounts_and_price_lists_persistence(self):
        data = {
            'name': 'Soriana V.D',
            'sku_code': 'SO1',
            'category': self.category.id,
            'retail_price': '75.00',
            'default_purchase_cost': '50.00',
            'has_volume_discounts': 'on',
            'volume_min_qty[]': ['10', '50'],
            'volume_disc_pct[]': ['5.00', '15.00'],
            'price_list_name[]': ['VIP Pharmacies', 'Wholesalers'],
            'price_list_price[]': ['70.00', '65.00'],
        }
        resp = self.client.post(reverse('product_create'), data, follow=True)
        self.assertEqual(resp.status_code, 200)
        p = Product.objects.get(sku_code='SO1')
        self.assertEqual(p.volume_discounts.count(), 2)
        self.assertEqual(p.price_lists.count(), 2)

        # Test detail view renders them
        detail_resp = self.client.get(reverse('product_detail', args=[p.id]))
        self.assertEqual(detail_resp.status_code, 200)
        self.assertContains(detail_resp, 'VIP Pharmacies')
        self.assertContains(detail_resp, '15,00%')
        self.assertTrue(p.volume_discounts.filter(discount_percentage=Decimal('15.00')).exists())

