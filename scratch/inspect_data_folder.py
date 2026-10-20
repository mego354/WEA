import os
import sys
sys.stdout.reconfigure(encoding='utf-8')
import openpyxl

data_folder = r'c:\D\Cooding\GENERAL\new\ERP\data_folder'
files = [f for f in os.listdir(data_folder) if f.endswith('.xlsx')]

out_path = r'c:\D\Cooding\GENERAL\new\ERP\scratch\data_inspection.txt'
with open(out_path, 'w', encoding='utf-8') as out:
    def log(msg):
        print(msg)
        out.write(str(msg) + '\n')

    log(f"Found {len(files)} excel files in {data_folder}")

    for f in sorted(files):
        path = os.path.join(data_folder, f)
        wb = openpyxl.load_workbook(path, read_only=True)
        log(f"\n==========================================")
        log(f"FILE: {f} (Sheets: {wb.sheetnames})")
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows = list(ws.iter_rows(values_only=True))
            row_count = len(rows)
            headers = rows[0] if row_count > 0 else []
            log(f"  Sheet '{sheet_name}': {row_count} total rows (including header)")
            log(f"  Headers: {headers}")
            if row_count > 1:
                log(f"  Sample row 1: {rows[1]}")
            if row_count > 2:
                log(f"  Sample row 2: {rows[2]}")
        wb.close()



