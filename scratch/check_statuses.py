import openpyxl
import sys
import io

if sys.platform == 'win32':
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

def check_statuses_and_dates():
    # 1. Sales Orders Statuses
    wb = openpyxl.load_workbook('data_folder/اوامر البيع.xlsx', data_only=True)
    ws = wb.active
    h = [str(c).strip() for c in next(ws.iter_rows(values_only=True))]
    status_idx = h.index('الحالة - الاسم')
    code_idx = h.index('الحالة - الكود')
    statuses = set()
    sample_dates = []
    date_idx = h.index('تاريخ الإنشاء')
    for r in list(ws.iter_rows(values_only=True))[1:1000]:
        statuses.add((r[code_idx], r[status_idx]))
        if len(sample_dates) < 5 and r[date_idx]:
            sample_dates.append((type(r[date_idx]), r[date_idx]))
    wb.close()
    print("Sales Order Statuses:", statuses)
    print("Sample Sales Order Dates:", sample_dates)

    # 2. Collections Statuses
    wb = openpyxl.load_workbook('data_folder/التحصيلات.xlsx', data_only=True)
    ws = wb.active
    h = [str(c).strip() for c in next(ws.iter_rows(values_only=True))]
    col_statuses = set()
    col_dates = []
    status_idx = h.index('اسم حالة التحصيل')
    date_idx = h.index('تاريخ الإنشاء')
    for r in list(ws.iter_rows(values_only=True))[1:1000]:
        col_statuses.add(r[status_idx])
        if len(col_dates) < 5 and r[date_idx]:
            col_dates.append((type(r[date_idx]), r[date_idx]))
    wb.close()
    print("Collection Statuses:", col_statuses)
    print("Sample Collection Dates:", col_dates)

    # 3. Customer Statuses
    wb = openpyxl.load_workbook('data_folder/لستة العملاء.xlsx', data_only=True)
    ws = wb.active
    h = [str(c).strip() for c in next(ws.iter_rows(values_only=True))]
    cust_statuses = set()
    status_idx = h.index('حالة العميل - الاسم')
    for r in list(ws.iter_rows(values_only=True))[1:1000]:
        cust_statuses.add(r[status_idx])
    wb.close()
    print("Customer Statuses:", cust_statuses)

if __name__ == '__main__':
    check_statuses_and_dates()
