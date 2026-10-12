import os
import sys
import openpyxl
import json
from decimal import Decimal
import datetime

if sys.platform == 'win32':
    import io
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

def deep_analyze():
    # 1. Customers
    print("=" * 60)
    print("1. ANALYZING: لستة العملاء.xlsx")
    wb = openpyxl.load_workbook('data_folder/لستة العملاء.xlsx', data_only=True)
    ws = wb.active
    headers = [str(c).strip() if c else '' for c in next(ws.iter_rows(values_only=True))]
    rows = list(ws.iter_rows(values_only=True))[1:]
    wb.close()
    
    print(f"Total Rows: {len(rows)}")
    id_idx = headers.index('المعرف')
    name_idx = headers.index('الاسم')
    serial_idx = headers.index('مسلسل')
    code_idx = headers.index('الكود')
    num_idx = headers.index('الرقم')
    area_idx = headers.index('المنطقة - الاسم')
    gov_idx = headers.index('المنطقة - المنطقة الرئيسية - الاسم')
    mobile_idx = headers.index('الموبايل')
    cat_idx = headers.index('تصنيف العميل')
    type_idx = headers.index('نوع الاكونت')
    
    ids = set()
    serials = set()
    names = set()
    for r in rows:
        if r[id_idx]: ids.add(r[id_idx])
        if r[serial_idx]: serials.add(r[serial_idx])
        if r[name_idx]: names.add(r[name_idx])
        
    print(f"Unique المعرف: {len(ids)}, Unique مسلسل: {len(serials)}, Unique الاسم: {len(names)}")
    print("Sample Row 0:")
    for h, v in zip(headers, rows[0]):
        if v is not None and str(v).strip():
            print(f"  {h}: {v}")

    # 2. Sales Orders & Products
    print("\n" + "=" * 60)
    print("2. ANALYZING: اوامر البيع.xlsx")
    wb = openpyxl.load_workbook('data_folder/اوامر البيع.xlsx', data_only=True)
    ws = wb.active
    so_headers = [str(c).strip() if c else '' for c in next(ws.iter_rows(values_only=True))]
    so_rows = list(ws.iter_rows(values_only=True))[1:]
    wb.close()
    
    print(f"Total Rows: {len(so_rows)}")
    so_id_idx = so_headers.index('المعرف')
    so_cust_idx = so_headers.index('العميل - المعرف')
    so_prod_name_idx = so_headers.index('العناصر - المنتج - الاسم')
    so_prod_sku_idx = so_headers.index('العناصر - المنتج - الكود')
    so_prod_cat_idx = so_headers.index('العناصر - المنتج - التصنيف')
    so_prod_barcode_idx = so_headers.index('العناصر - المنتج - باركود')
    so_prod_price_idx = so_headers.index('العناصر - المنتج - سعر القطعة')
    
    order_ids = set()
    products = {} # sku: info
    categories = set()
    
    for r in so_rows:
        order_ids.add(r[so_id_idx])
        p_name = r[so_prod_name_idx]
        p_sku = r[so_prod_sku_idx]
        p_cat = r[so_prod_cat_idx]
        p_barcode = r[so_prod_barcode_idx]
        p_price = r[so_prod_price_idx]
        if p_sku:
            products[p_sku] = {
                'name': p_name,
                'category': p_cat,
                'barcode': p_barcode,
                'retail_price': p_price
            }
        if p_cat:
            categories.add(p_cat)
            
    print(f"Unique Orders: {len(order_ids)}")
    print(f"Unique Products discovered: {len(products)}")
    print(f"Unique Product Categories discovered: {len(categories)}")
    print("Categories:", list(categories))
    print("Sample Products (first 5):")
    for sku, pinfo in list(products.items())[:5]:
        print(f"  SKU: {sku} | {pinfo['name']} | Cat: {pinfo['category']} | Price: {pinfo['retail_price']}")

    # 3. Collections
    print("\n" + "=" * 60)
    print("3. ANALYZING: التحصيلات.xlsx")
    wb = openpyxl.load_workbook('data_folder/التحصيلات.xlsx', data_only=True)
    ws = wb.active
    col_headers = [str(c).strip() if c else '' for c in next(ws.iter_rows(values_only=True))]
    col_rows = list(ws.iter_rows(values_only=True))[1:]
    wb.close()
    print(f"Total Rows: {len(col_rows)}")
    print("Headers:", col_headers)
    print("Sample Row 0:")
    for h, v in zip(col_headers, col_rows[0]):
        if v is not None and str(v).strip():
            print(f"  {h}: {v}")

    # 4. Agreements
    print("\n" + "=" * 60)
    print("4. ANALYZING: الاتفاقيات.xlsx")
    wb = openpyxl.load_workbook('data_folder/الاتفاقيات.xlsx', data_only=True)
    ws = wb.active
    agr_headers = [str(c).strip() if c else '' for c in next(ws.iter_rows(values_only=True))]
    agr_rows = list(ws.iter_rows(values_only=True))[1:]
    wb.close()
    print(f"Total Rows: {len(agr_rows)}")
    print("Headers:", agr_headers)
    print("Sample Row 0:")
    for h, v in zip(agr_headers, agr_rows[0]):
        if v is not None and str(v).strip():
            print(f"  {h}: {v}")

    # 5. Returns
    print("\n" + "=" * 60)
    print("5. ANALYZING: المرتجعات.xlsx")
    wb = openpyxl.load_workbook('data_folder/المرتجعات.xlsx', data_only=True)
    ws = wb.active
    ret_headers = [str(c).strip() if c else '' for c in next(ws.iter_rows(values_only=True))]
    ret_rows = list(ws.iter_rows(values_only=True))[1:]
    wb.close()
    print(f"Total Rows: {len(ret_rows)}")
    print("Sample Row 0:")
    for h, v in zip(ret_headers, ret_rows[0]):
        if v is not None and str(v).strip():
            print(f"  {h}: {v}")

if __name__ == '__main__':
    deep_analyze()
