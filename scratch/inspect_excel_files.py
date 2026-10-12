import os
import sys
import openpyxl
import json
from decimal import Decimal
import datetime

if sys.platform == 'win32':
    import io
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

DATA_FOLDER = "data_folder"

def serialize(obj):
    if isinstance(obj, (datetime.datetime, datetime.date, datetime.time)):
        return obj.isoformat()
    if isinstance(obj, Decimal):
        return str(obj)
    return str(obj)

def inspect_all():
    files = [f for f in os.listdir(DATA_FOLDER) if f.endswith(('.xlsx', '.xls'))]
    all_info = {}
    
    for filename in sorted(files):
        filepath = os.path.join(DATA_FOLDER, filename)
        wb = openpyxl.load_workbook(filepath, data_only=True)
        all_info[filename] = {
            "size_bytes": os.path.getsize(filepath),
            "sheets": {}
        }
        
        for sheetname in wb.sheetnames:
            ws = wb[sheetname]
            all_rows = list(ws.iter_rows(values_only=True))
            if not all_rows:
                continue
                
            raw_headers = all_rows[0]
            headers = [str(c).strip() if c is not None else f"Col_{i}" for i, c in enumerate(raw_headers)]
            data_rows = all_rows[1:]
            
            # Count actual non-empty rows
            non_empty_rows = [r for r in data_rows if any(cell is not None and str(cell).strip() != '' for cell in r)]
            
            # Null counts per column
            null_counts = {h: 0 for h in headers}
            for r in non_empty_rows:
                for idx, h in enumerate(headers):
                    val = r[idx] if idx < len(r) else None
                    if val is None or str(val).strip() == '':
                        null_counts[h] += 1
                        
            samples = []
            for r in non_empty_rows[:3]:
                row_dict = {}
                for idx, h in enumerate(headers):
                    val = r[idx] if idx < len(r) else None
                    row_dict[h] = serialize(val) if val is not None else None
                samples.append(row_dict)
                
            all_info[filename]["sheets"][sheetname] = {
                "max_row": ws.max_row,
                "max_col": ws.max_column,
                "total_data_rows": len(non_empty_rows),
                "columns": headers,
                "null_counts": null_counts,
                "sample_rows": samples
            }
        wb.close()
        
    with open("scratch/excel_inspection.json", "w", encoding="utf-8") as f:
        json.dump(all_info, f, ensure_ascii=False, indent=2)
        
    print(f"Successfully inspected {len(files)} files.")
    for fname, finfo in all_info.items():
        print("=" * 60)
        print(f"File: {fname} ({finfo['size_bytes']:,} bytes)")
        for sname, sinfo in finfo["sheets"].items():
            print(f"  Sheet: {sname} -> Rows: {sinfo['total_data_rows']:,}, Cols: {len(sinfo['columns'])}")
            print(f"  Columns: {sinfo['columns']}")

if __name__ == "__main__":
    inspect_all()
